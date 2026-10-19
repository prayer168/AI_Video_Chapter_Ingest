#!/usr/bin/env python3
"""Build accessible HTML and print-ready DOCX/PDF outputs from Markdown."""
from __future__ import annotations

import argparse
import html
import os
import re
import shutil
import subprocess
import sys
from pathlib import Path
from urllib.parse import urlparse

import markdown
import yaml
from image_policy import (
    IMAGE_SUFFIXES, approved_image_paths, publication_blockers,
    redact_unapproved_images,
)

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "dist"
HTML = OUT / "html"
IMAGE_PATTERN = re.compile(r"(!\[[^\]]*\]\()([^\s)]+)([^)]*\))")


def source_files() -> list[Path]:
    paths = [ROOT / "README.md", ROOT / "docs" / "TOC.md"]
    paths += sorted((ROOT / "docs").glob("*.md"))
    paths += sorted((ROOT / "chapters").glob("*.md"))
    return [path for path in paths if path.exists()]


def book_chapters() -> list[Path]:
    """Return only canonical chapter manuscripts for reader-edition exports."""
    return sorted(
        path
        for path in (ROOT / "chapters").glob("*.md")
        if path.name[:2].isdigit()
    )


def split_front_matter(text: str) -> tuple[dict, str]:
    match = re.match(r"\A---\s*\n(.*?)\n---\s*\n", text, re.S)
    if not match:
        return {}, text
    try:
        metadata = yaml.safe_load(match.group(1)) or {}
    except yaml.YAMLError:
        metadata = {}
    return metadata, text[match.end():]


def output_path(source: Path) -> Path:
    if source == ROOT / "README.md":
        return HTML / "index.html"
    if source == ROOT / "docs" / "TOC.md":
        return HTML / "toc.html"
    if source.parent == ROOT / "docs":
        return HTML / "docs" / f"{source.stem}.html"
    return HTML / f"{source.stem}.html"


def rewrite_links(text: str, source: Path, target: Path, known_sources: set[Path]) -> str:
    pattern = re.compile(r"(!?\[[^\]]*\]\()([^)]*)(\))")

    def replace(match: re.Match) -> str:
        label, raw_target, close = match.groups()
        target_text = raw_target.strip()
        if not target_text or target_text.startswith("#"):
            return match.group(0)
        parsed = urlparse(target_text)
        if parsed.scheme or parsed.netloc:
            return match.group(0)
        file_part, separator, fragment = target_text.partition("#")
        resolved = (source.parent / file_part).resolve()
        if resolved in known_sources:
            destination = output_path(resolved)
        elif resolved.is_relative_to(ROOT / "assets"):
            destination = HTML / "assets" / resolved.relative_to(ROOT / "assets")
        elif resolved.is_relative_to(ROOT / "chapters"):
            relative = resolved.relative_to(ROOT / "chapters")
            if relative.parts and relative.parts[0].endswith("-assets"):
                destination = HTML / relative
            else:
                return match.group(0)
        else:
            return match.group(0)
        relative_url = Path(os.path.relpath(destination, target.parent)).as_posix()
        if separator:
            relative_url += "#" + fragment
        return f"{label}{relative_url}{close}"

    return pattern.sub(replace, text)


def markdown_body(text: str, source: Path, target: Path, known_sources: set[Path]) -> str:
    _, body = split_front_matter(text)
    body = rewrite_links(body, source, target, known_sources)
    return markdown.markdown(
        body,
        extensions=["tables", "fenced_code", "toc", "sane_lists"],
        output_format="html5",
    )


def page_title(source: Path, text: str) -> str:
    metadata, body = split_front_matter(text)
    if metadata.get("title"):
        return str(metadata["title"])
    heading = re.search(r"^#\s+(.+)$", body, re.M)
    if heading:
        return heading.group(1).strip()
    if source.name == "README.md":
        return "AI 電影製作聖經"
    return source.stem.replace("-", " ").title()


def build_html(sources: list[Path], *, preview: bool = False, publication_target: str = "zh-TW-site", approved: set[Path] | None = None) -> None:
    approved = approved or set()
    HTML.mkdir(parents=True, exist_ok=True)
    known = {path.resolve() for path in sources}
    for source in sources:
        target_path = output_path(source)
        target_path.parent.mkdir(parents=True, exist_ok=True)
        text = source.read_text(encoding="utf-8-sig")
        if preview:
            text = redact_unapproved_images(text, ROOT, source, publication_target)
        title = page_title(source, text)
        content = markdown_body(text, source, target_path, known)
        if preview:
            content = '<p class="draft-preview-notice"><strong>草稿預覽</strong>：未核准的插圖已省略，正式出版前須完成審查。</p>' + content
        page = f"""<!doctype html>
<html lang="zh-Hant">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<meta name="description" content="AI 電影製作聖經章節與出版資料">
<title>{html.escape(title)}</title>
<style>
:root{{color-scheme:light;--ink:#183247;--muted:#526777;--line:#dce5eb;--accent:#176b87}}
*{{box-sizing:border-box}}body{{margin:0;color:var(--ink);font:18px/1.8 system-ui,"Noto Sans CJK TC","Microsoft JhengHei",sans-serif}}
main{{max-width:900px;margin:0 auto;padding:2rem clamp(1rem,4vw,3rem) 5rem}}
h1,h2,h3{{line-height:1.35;margin-top:2.2em}}h1{{font-size:2.2rem;margin-top:1em}}h2{{border-bottom:1px solid var(--line);padding-bottom:.3em}}
a{{color:var(--accent);overflow-wrap:anywhere}}img,svg{{max-width:100%;height:auto}}figure{{margin:2rem 0}}figcaption,blockquote{{color:var(--muted)}}
blockquote{{border-left:4px solid var(--accent);margin:1.5rem 0;padding:.2rem 1rem}}table{{border-collapse:collapse;display:block;max-width:100%;overflow-x:auto;font-size:.95rem}}th,td{{border:1px solid var(--line);padding:.45rem .65rem;text-align:left;vertical-align:top}}th{{background:#edf4f7}}
pre{{overflow:auto;padding:1rem;background:#f2f5f7;border-radius:8px}}code{{font-family:ui-monospace,Consolas,monospace;font-size:.9em}}hr{{border:0;border-top:1px solid var(--line);margin:2rem 0}}
@media(max-width:600px){{body{{font-size:16px}}main{{padding:1rem 1rem 3rem}}h1{{font-size:1.8rem}}}}
@media print{{body{{font-size:11pt}}main{{max-width:none;padding:0}}a{{color:inherit;text-decoration:none}}}}
</style>
</head>
<body><main class="markdown-body">{content}</main></body>
</html>"""
        target_path.write_text(page, encoding="utf-8")
    def ignore_unapproved_images(directory: str, names: list[str]) -> set[str]:
        base = Path(directory)
        return {
            name for name in names
            if (base / name).suffix.casefold() in IMAGE_SUFFIXES
            and (base / name).resolve() not in approved
        }

    assets = ROOT / "assets"
    if assets.exists():
        shutil.copytree(assets, HTML / "assets", dirs_exist_ok=True, ignore=ignore_unapproved_images)
    for folder in (ROOT / "chapters").glob("*-assets"):
        shutil.copytree(folder, HTML / folder.name, dirs_exist_ok=True, ignore=ignore_unapproved_images)
    links = []
    for source in sources:
        destination = output_path(source).relative_to(HTML).as_posix()
        links.append(f'<li><a href="{html.escape(destination)}">{html.escape(page_title(source, source.read_text(encoding="utf-8-sig")))}</a></li>')
    index = "<!doctype html><html lang='zh-Hant'><meta charset='utf-8'><meta name='viewport' content='width=device-width'><title>書籍內容</title><h1>AI 電影製作聖經</h1><nav><ul>" + "".join(links) + "</ul></nav></html>"
    (HTML / "contents.html").write_text(index, encoding="utf-8")
    print(f"Built HTML: {HTML}")


def combined_markdown(sources: list[Path], *, preview: bool = False, target: str = "zh-TW-book") -> Path:
    sections = []
    for source in sources:
        text = source.read_text(encoding="utf-8-sig")
        metadata, body = split_front_matter(text)
        title = metadata.get("title") or page_title(source, text)
        if preview:
            body = redact_unapproved_images(body, ROOT, source, target)
        body = rewrite_image_paths(body, source.parent)
        # Chapter manuscripts already carry their own H1. Avoid adding a
        # duplicate title that would create two chapter headings in exports.
        if not re.search(r"^#\s+", body, re.M):
            body = f"# {title}\n\n{body}"
        sections.append(body.strip())
    target = OUT / "book-source.md"
    prefix = "> 草稿預覽：未通過出版審查的插圖已省略。\n\n" if preview else ""
    target.write_text(prefix + "\n\n\\newpage\n\n".join(sections), encoding="utf-8")
    return target


def rewrite_image_paths(text: str, base_dir: Path) -> str:
    """Rewrite existing local Markdown images to paths rooted at the repo."""
    def replace(match: re.Match) -> str:
        prefix, raw_target, suffix = match.groups()
        parsed = urlparse(raw_target)
        if parsed.scheme or parsed.netloc or raw_target.startswith("#"):
            return match.group(0)
        image_path = (base_dir / raw_target).resolve()
        if not image_path.exists() or not image_path.is_relative_to(ROOT):
            return match.group(0)
        rooted_path = image_path.relative_to(ROOT).as_posix()
        return f"{prefix}{rooted_path}{suffix}"

    return IMAGE_PATTERN.sub(replace, text)


def convert_svg_images(source: Path, output_format: str) -> int:
    """Convert SVGs to a format supported by the selected Pandoc writer."""
    svg_converter = shutil.which("inkscape")
    media_dir = OUT / f"{output_format}-media"
    source_text = source.read_text(encoding="utf-8")

    def convert(match: re.Match) -> str:
        prefix, raw_target, suffix = match.groups()
        if Path(urlparse(raw_target).path).suffix.casefold() != ".svg":
            return match.group(0)
        image_path = (ROOT / raw_target).resolve()
        if not image_path.exists():
            raise FileNotFoundError(f"SVG image not found: {image_path}")
        if not svg_converter:
            raise RuntimeError("This export contains SVG figures. Install Inkscape and ensure 'inkscape' is on PATH.")
        media_dir.mkdir(parents=True, exist_ok=True)
        relative_name = "--".join(image_path.relative_to(ROOT).parts).removesuffix(".svg")
        converted = media_dir / f"{relative_name}.{output_format}"
        result = subprocess.run(
            [svg_converter, str(image_path), f"--export-filename={converted}"],
            cwd=ROOT,
        )
        if result.returncode:
            raise RuntimeError(f"Inkscape failed to convert {image_path.name} to {output_format.upper()}")
        rooted_path = converted.relative_to(ROOT).as_posix()
        return f"{prefix}{rooted_path}{suffix}"

    try:
        source_text = IMAGE_PATTERN.sub(convert, source_text)
    except (FileNotFoundError, RuntimeError) as exc:
        print(str(exc), file=sys.stderr)
        return 2
    source.write_text(source_text, encoding="utf-8")
    return 0


def build_docx(sources: list[Path], font: str, *, preview: bool = False, target_scope: str = "zh-TW-book") -> int:
    pandoc = shutil.which("pandoc")
    source = combined_markdown(book_chapters(), preview=preview, target=target_scope)
    target = OUT / "AI-Filmmaking-Bible.docx"
    if pandoc:
        conversion_status = convert_svg_images(source, "png")
        if conversion_status:
            return conversion_status
        result = subprocess.run(
            [pandoc, str(source), "-o", str(target), "-V", f"mainfont={font}"],
            cwd=ROOT,
        )
        if result.returncode:
            return result.returncode
        try:
            from docx import Document
            from docx.oxml.ns import qn
        except ImportError:
            print("DOCX created; install python-docx to apply the requested CJK font to Word styles.", file=sys.stderr)
        else:
            document = Document(target)
            for style_name in ("Normal", "Heading 1", "Heading 2", "Heading 3"):
                style = document.styles[style_name]
                style.font.name = font
                properties = style._element.get_or_add_rPr()
                fonts = properties.rFonts
                if fonts is None:
                    from docx.oxml import OxmlElement
                    fonts = OxmlElement("w:rFonts")
                    properties.append(fonts)
                for key in ("ascii", "hAnsi", "eastAsia", "cs"):
                    fonts.set(qn(f"w:{key}"), font)
            document.save(target)
        print(f"Built DOCX{' draft preview' if preview else ''}: {target}")
        return 0
    try:
        from docx import Document
    except ImportError:
        print("DOCX requires Pandoc or python-docx: pip install python-docx", file=sys.stderr)
        return 2
    document = Document()
    document.styles["Normal"].font.name = font
    for source_path in sources:
        text = source_path.read_text(encoding="utf-8-sig")
        _, body = split_front_matter(text)
        for line in body.splitlines():
            stripped = line.strip()
            if not stripped:
                continue
            if stripped.startswith("#"):
                level = min(len(stripped) - len(stripped.lstrip("#")), 3)
                document.add_heading(stripped.lstrip("# ").strip(), level)
            elif stripped.startswith("- "):
                document.add_paragraph(stripped[2:], style="List Bullet")
            else:
                document.add_paragraph(stripped)
    document.save(target)
    print(f"Built DOCX{' draft preview' if preview else ''}: {target}")
    return 0


def build_pdf(sources: list[Path], font: str, *, preview: bool = False, target_scope: str = "zh-TW-book") -> int:
    pandoc, xelatex = shutil.which("pandoc"), shutil.which("xelatex")
    if not pandoc or not xelatex:
        print("PDF requires Pandoc and XeLaTeX on PATH", file=sys.stderr)
        return 2
    source = combined_markdown(book_chapters(), preview=preview, target=target_scope)
    conversion_status = convert_svg_images(source, "pdf")
    if conversion_status:
        return conversion_status
    target = OUT / "AI-Filmmaking-Bible.pdf"
    result = subprocess.run(
        [
            pandoc,
            str(source),
            "-o",
            str(target),
            "--pdf-engine=xelatex",
            "-V",
            f"mainfont={font}",
            "-V",
            f"CJKmainfont={font}",
        ],
        cwd=ROOT,
    )
    if result.returncode == 0:
        print(f"Built PDF{' draft preview' if preview else ''}: {target}")
    return result.returncode



def build_epub(font: str = "Noto Serif CJK TC", *, preview: bool = False, target_scope: str = "zh-TW-book") -> int:
    pandoc = shutil.which("pandoc")
    if not pandoc:
        print("EPUB requires Pandoc on PATH", file=sys.stderr)
        return 2
    source = combined_markdown(book_chapters(), preview=preview, target=target_scope)
    conversion_status = convert_svg_images(source, "png")
    if conversion_status:
        return conversion_status
    target = OUT / "AI-Filmmaking-Bible.epub"
    result = subprocess.run(
        [pandoc, str(source), "-o", str(target), "--split-level=1", "-M", "title=AI 電影製作聖經", "-M", "lang=zh-TW"],
        cwd=ROOT,
    )
    if result.returncode == 0:
        print(f"Built EPUB{' draft preview' if preview else ''}: {target}")
    return result.returncode


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--format", required=True, choices=["html", "docx", "pdf", "epub", "all"])
    parser.add_argument("--preview", action="store_true", help="Build a draft preview and replace unapproved images with notices")
    parser.add_argument("--font", default="Noto Serif CJK TC")
    args = parser.parse_args()
    formats = ["html", "docx", "pdf", "epub"] if args.format == "all" else [args.format]
    def scope_for(output_format: str) -> str:
        return "zh-TW-site" if output_format == "html" else "zh-TW-book"

    if args.preview:
        global OUT, HTML
        OUT = ROOT / "dist" / "preview"
        HTML = OUT / "html"
    else:
        all_blockers: list[str] = []
        for output_format in formats:
            scope = scope_for(output_format)
            blockers = publication_blockers(ROOT, scope, include_docs=output_format == "html")
            all_blockers.extend(f"{scope}: {item}" for item in blockers)
        if all_blockers:
            print("PUBLICATION BLOCKED: resolve image review items or use --preview", file=sys.stderr)
            for blocker in dict.fromkeys(all_blockers):
                print(f"- {blocker}", file=sys.stderr)
            return 2
    sources = source_files()
    OUT.mkdir(parents=True, exist_ok=True)
    for output_format in formats:
        target_scope = scope_for(output_format)
        approved = approved_image_paths(ROOT, target_scope)
        if output_format == "html":
            build_html(sources, preview=args.preview, publication_target=target_scope, approved=approved)
        elif output_format == "docx":
            result = build_docx(sources, args.font, preview=args.preview, target_scope=target_scope)
            if result:
                return result
        elif output_format == "pdf":
            result = build_pdf(sources, args.font, preview=args.preview, target_scope=target_scope)
            if result:
                return result
        else:
            result = build_epub(args.font, preview=args.preview, target_scope=target_scope)
            if result:
                return result
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
